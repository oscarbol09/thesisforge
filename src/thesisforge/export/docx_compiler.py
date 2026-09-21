"""APA 7th Edition compliant DOCX compiler for academic thesis projects."""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from thesisforge.drafting.sanitizer import sanitize_cell_value
from thesisforge.exceptions import ExportError
from thesisforge.models import CitationDTO, ExportOptionsDTO, ProjectStateDTO

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph


class APA7DocxCompiler:
    """Compiles structured thesis projects into strictly formatted APA 7th edition Word documents."""

    def __init__(self, default_options: ExportOptionsDTO | None = None) -> None:
        self.default_options = default_options or ExportOptionsDTO()

    def compile(
        self,
        project: ProjectStateDTO,
        options: ExportOptionsDTO | None = None,
    ) -> DocxDocument:
        """Compile a full ProjectStateDTO into an APA 7 docx.Document object."""
        opts = options or self.default_options
        doc: DocxDocument = Document()

        # 1. Page Geometry & Default Margins (1.0 inch / 2.54 cm)
        self._setup_page_layout(doc, opts)

        # 2. Typography Styles
        self._setup_styles(doc, opts)

        # 3. Document Metadata
        self._set_core_metadata(doc, project, opts)

        # 4. Header with Page Numbering
        self._setup_headers(doc, opts)

        # 5. Title / Cover Page (if enabled)
        if opts.include_cover_page:
            self._render_cover_page(doc, project, opts)

        # 6. Body Sections (Sorted by chapter_number and order_index)
        sorted_sections = sorted(
            project.sections,
            key=lambda s: (s.chapter_number, s.order_index),
        )

        current_chapter: int | None = None
        for section in sorted_sections:
            if current_chapter is not None and section.chapter_number != current_chapter:
                # Chapter separator / page break between major chapters
                doc.add_page_break()  # type: ignore[no-untyped-call]

            current_chapter = section.chapter_number
            self._render_section(doc, section.title, section.content, opts)

        # 7. References Section (if enabled)
        if opts.include_references and project.validated_citations:
            doc.add_page_break()  # type: ignore[no-untyped-call]
            self._render_references(doc, project.validated_citations, opts)

        return doc

    def compile_to_bytes(
        self,
        project: ProjectStateDTO,
        options: ExportOptionsDTO | None = None,
    ) -> bytes:
        """Compile project and return raw bytes of the DOCX file."""
        try:
            doc = self.compile(project, options)
            buffer = io.BytesIO()
            doc.save(buffer)
            return buffer.getvalue()
        except Exception as e:
            if isinstance(e, ExportError):
                raise
            raise ExportError(f"Error al compilar documento Word APA 7: {e}") from e

    def compile_to_file(
        self,
        project: ProjectStateDTO,
        file_path: str | Path,
        options: ExportOptionsDTO | None = None,
    ) -> Path:
        """Compile project and save to designated file system path."""
        target = Path(file_path)
        target.parent.mkdir(parents=True, exist_ok=True)
        docx_bytes = self.compile_to_bytes(project, options)
        target.write_bytes(docx_bytes)
        return target

    def _setup_page_layout(self, doc: DocxDocument, options: ExportOptionsDTO) -> None:
        """Apply 1-inch (2.54 cm) margins on all sides to all sections."""
        for section in doc.sections:
            margin = Inches(options.margin_inches)
            section.top_margin = margin
            section.bottom_margin = margin
            section.left_margin = margin
            section.right_margin = margin
            section.page_width = Inches(8.5)
            section.page_height = Inches(11.0)

    def _setup_styles(self, doc: DocxDocument, options: ExportOptionsDTO) -> None:
        """Configure font and line spacing across default styles."""
        normal_style = doc.styles["Normal"]
        font = normal_style.font
        font.name = options.font_name
        font.size = Pt(options.font_size_pt)
        font.color.rgb = RGBColor(0, 0, 0)

        p_format = normal_style.paragraph_format
        p_format.line_spacing = options.line_spacing
        p_format.space_before = Pt(0)
        p_format.space_after = Pt(0)
        p_format.first_line_indent = Inches(0.5)

    def _set_core_metadata(
        self,
        doc: DocxDocument,
        project: ProjectStateDTO,
        options: ExportOptionsDTO,
    ) -> None:
        """Sanitize and set document core properties cleanly."""
        props = doc.core_properties
        props.title = project.title or "Tesis de Grado"
        props.author = options.author_name or "Investigador"
        props.subject = project.area_of_study or "Investigación Académica"
        props.comments = "Documento académico estructurado bajo norma APA 7ma edición."
        props.category = "Tesis de Grado"
        props.language = project.language or "es-ES"

    def _setup_headers(self, doc: DocxDocument, options: ExportOptionsDTO) -> None:
        """Configure APA 7 running header with right-aligned page number."""
        for section in doc.sections:
            header = section.header
            p = header.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            run = p.add_run()
            font = run.font
            font.name = options.font_name
            font.size = Pt(options.font_size_pt)
            font.color.rgb = RGBColor(0, 0, 0)

            ns = nsdecls("w")
            fld = parse_xml(f'<w:fldSimple {ns} w:instr="PAGE"/>')
            run._r.append(fld)

    def _render_cover_page(
        self,
        doc: DocxDocument,
        project: ProjectStateDTO,
        options: ExportOptionsDTO,
    ) -> None:
        """Render standard APA 7 Student/Professional Title Page."""
        for _ in range(3):
            blank_p = doc.add_paragraph()
            blank_p.paragraph_format.first_line_indent = Inches(0)
            blank_p.paragraph_format.line_spacing = options.line_spacing
            blank_p.paragraph_format.space_before = Pt(0)
            blank_p.paragraph_format.space_after = Pt(0)

        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_p.paragraph_format.first_line_indent = Inches(0)
        title_p.paragraph_format.line_spacing = options.line_spacing
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(0)
        title_run = title_p.add_run(project.title or "Título de la Investigación")
        title_run.bold = True
        title_run.font.name = options.font_name
        title_run.font.size = Pt(options.font_size_pt)

        sep_p = doc.add_paragraph()
        sep_p.paragraph_format.first_line_indent = Inches(0)
        sep_p.paragraph_format.line_spacing = options.line_spacing
        sep_p.paragraph_format.space_before = Pt(0)
        sep_p.paragraph_format.space_after = Pt(0)

        meta_lines: list[str] = []
        if options.author_name:
            meta_lines.append(options.author_name)
        if options.faculty_or_program:
            meta_lines.append(options.faculty_or_program)
        if options.institution_name:
            meta_lines.append(options.institution_name)
        if options.advisor_name:
            meta_lines.append(f"Asesor: {options.advisor_name}")
        if options.city_and_country and options.year:
            meta_lines.append(f"{options.city_and_country}, {options.year}")
        elif options.year:
            meta_lines.append(str(options.year))
        elif options.city_and_country:
            meta_lines.append(options.city_and_country)

        for line in meta_lines:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.line_spacing = options.line_spacing
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line)
            run.font.name = options.font_name
            run.font.size = Pt(options.font_size_pt)

        doc.add_page_break()  # type: ignore[no-untyped-call]

    def _render_section(
        self,
        doc: DocxDocument,
        section_title: str,
        content_md: str,
        options: ExportOptionsDTO,
    ) -> None:
        """Render a section draft with markdown parsing for headings, tables, and paragraphs."""
        self._add_heading(doc, section_title, level=1, options=options)

        if not content_md or not content_md.strip():
            return

        lines = content_md.splitlines()
        idx = 0
        n = len(lines)

        while idx < n:
            line = lines[idx]
            stripped = line.strip()

            if not stripped:
                idx += 1
                continue

            if stripped.startswith("#"):
                match = re.match(r"^(#{1,5})\s+(.+)$", stripped)
                if match:
                    h_level = len(match.group(1))
                    h_text = match.group(2)
                    self._add_heading(doc, h_text, level=min(h_level + 1, 5), options=options)
                    idx += 1
                    continue

            if stripped.startswith("|") and stripped.endswith("|"):
                table_lines: list[str] = []
                while (
                    idx < n
                    and lines[idx].strip().startswith("|")
                    and lines[idx].strip().endswith("|")
                ):
                    table_lines.append(lines[idx].strip())
                    idx += 1
                self._render_table(doc, table_lines, options)
                continue

            if stripped.startswith(">"):
                quote_lines: list[str] = []
                while idx < n and lines[idx].strip().startswith(">"):
                    quote_lines.append(re.sub(r"^>\s*", "", lines[idx].strip()))
                    idx += 1
                self._render_blockquote(doc, " ".join(quote_lines), options)
                continue

            if re.match(r"^[-*]\s+", stripped):
                list_text = re.sub(r"^[-*]\s+", "", stripped)
                self._render_list_item(doc, list_text, is_ordered=False, options=options)
                idx += 1
                continue

            num_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
            if num_match:
                list_text = num_match.group(2)
                self._render_list_item(doc, list_text, is_ordered=True, options=options)
                idx += 1
                continue

            para_lines: list[str] = [stripped]
            idx += 1
            while (
                idx < n
                and lines[idx].strip()
                and not lines[idx].strip().startswith(("#", "|", ">", "-", "*"))
                and not re.match(r"^\d+\.\s+", lines[idx].strip())
            ):
                para_lines.append(lines[idx].strip())
                idx += 1

            self._render_paragraph(doc, " ".join(para_lines), options)

    def _add_heading(
        self,
        doc: DocxDocument,
        text: str,
        level: int,
        options: ExportOptionsDTO,
    ) -> None:
        """Format 5 APA 7th Edition Heading Levels."""
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = options.line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Inches(0)
            run = p.add_run(text)
            run.bold = True
        elif level == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Inches(0)
            run = p.add_run(text)
            run.bold = True
        elif level == 3:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Inches(0)
            run = p.add_run(text)
            run.bold = True
            run.italic = True
        elif level == 4:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Inches(0.5)
            h_text = text if text.endswith((".", ":", "?")) else f"{text}."
            run = p.add_run(h_text)
            run.bold = True
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Inches(0.5)
            h_text = text if text.endswith((".", ":", "?")) else f"{text}."
            run = p.add_run(h_text)
            run.bold = True
            run.italic = True

        run.font.name = options.font_name
        run.font.size = Pt(options.font_size_pt)
        run.font.color.rgb = RGBColor(0, 0, 0)

    def _render_paragraph(
        self,
        doc: DocxDocument,
        text: str,
        options: ExportOptionsDTO,
    ) -> None:
        """Render a standard body paragraph with 0.5 in first line indent and inline markdown formatting."""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.line_spacing = options.line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        self._add_inline_formatted_runs(p, text, options)

    def _render_blockquote(
        self,
        doc: DocxDocument,
        text: str,
        options: ExportOptionsDTO,
    ) -> None:
        """Render an APA 7 blockquote (0.5 in left indent, no first line indent, double-spaced)."""
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = options.line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        self._add_inline_formatted_runs(p, text, options)

    def _render_list_item(
        self,
        doc: DocxDocument,
        text: str,
        is_ordered: bool,
        options: ExportOptionsDTO,
    ) -> None:
        """Render an indented list item with bullet or number."""
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.line_spacing = options.line_spacing
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)

        bullet_prefix = "•  " if not is_ordered else "1.  "
        p.add_run(bullet_prefix)
        self._add_inline_formatted_runs(p, text, options)

    def _render_table(
        self,
        doc: DocxDocument,
        table_lines: list[str],
        options: ExportOptionsDTO,
    ) -> None:
        """Render an APA 7 styled table with formula sanitization and no vertical borders."""
        parsed_rows: list[list[str]] = []
        for line in table_lines:
            stripped = line.strip()
            if not stripped:
                continue
            inner = stripped
            if inner.startswith("|"):
                inner = inner[1:]
            if inner.endswith("|"):
                inner = inner[:-1]

            raw_cells = re.split(r"(?<!\\)\|", inner)
            cells = [c.replace(r"\|", "|").strip() for c in raw_cells]
            if all(re.match(r"^:?-+:?$", c) for c in cells if c):
                continue
            parsed_rows.append(cells)

        if not parsed_rows:
            return

        num_cols = max(len(row) for row in parsed_rows)
        table: Table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for row_idx, row_data in enumerate(parsed_rows):
            for col_idx in range(num_cols):
                raw_val = row_data[col_idx] if col_idx < len(row_data) else ""
                sanitized_val = sanitize_cell_value(raw_val)

                cell = table.cell(row_idx, col_idx)
                cell_p = cell.paragraphs[0]
                cell_p.paragraph_format.first_line_indent = Inches(0)
                cell_p.paragraph_format.line_spacing = 1.15
                cell_p.paragraph_format.space_before = Pt(2)
                cell_p.paragraph_format.space_after = Pt(2)

                run = cell_p.add_run(sanitized_val)
                run.font.name = options.font_name
                run.font.size = Pt(options.font_size_pt - 1)
                run.font.color.rgb = RGBColor(0, 0, 0)
                if row_idx == 0:
                    run.bold = True

        self._apply_apa7_table_borders(table)

        spacer = doc.add_paragraph()
        spacer.paragraph_format.first_line_indent = Inches(0)
        spacer.paragraph_format.line_spacing = options.line_spacing
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(0)

    def _apply_apa7_table_borders(self, table: Table) -> None:
        """Apply APA 7 horizontal-only borders using Word XML."""
        tblPr = table._tbl.tblPr
        ns = nsdecls("w")
        tblBorders = parse_xml(
            f"<w:tblBorders {ns}>"
            r'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            r'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
            r'  <w:left w:val="none"/>'
            r'  <w:right w:val="none"/>'
            r'  <w:insideH w:val="none"/>'
            r'  <w:insideV w:val="none"/>'
            r"</w:tblBorders>"
        )
        tblPr.append(tblBorders)

        if len(table.rows) > 0:
            for cell in table.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(
                    f"<w:tcBorders {ns}>"
                    r'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
                    r"</w:tcBorders>"
                )
                tcPr.append(tcBorders)

    def _render_references(
        self,
        doc: DocxDocument,
        citations: list[CitationDTO],
        options: ExportOptionsDTO,
    ) -> None:
        """Render APA 7 References section with 0.5 in hanging indent."""
        self._add_heading(doc, "Referencias", level=1, options=options)

        def sort_key(c: CitationDTO) -> str:
            if c.authors:
                return c.authors[0].lower()
            return c.title.lower()

        sorted_citations = sorted(citations, key=sort_key)

        for citation in sorted_citations:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)  # Hanging indent
            p.paragraph_format.line_spacing = options.line_spacing
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            if citation.apa_formatted and citation.apa_formatted.strip():
                self._add_inline_formatted_runs(p, citation.apa_formatted, options)
            else:
                self._render_synthesized_apa_reference(p, citation, options)

    def _render_synthesized_apa_reference(
        self,
        p: Paragraph,
        citation: CitationDTO,
        options: ExportOptionsDTO,
    ) -> None:
        """Synthesize APA 7 reference format if pre-formatted string is absent."""
        if citation.authors:
            if len(citation.authors) == 1:
                auth_str = f"{citation.authors[0]} "
            elif len(citation.authors) == 2:
                auth_str = f"{citation.authors[0]} & {citation.authors[1]} "
            else:
                auth_str = f"{citation.authors[0]} et al. "
        else:
            auth_str = "Autor Desconocido. "

        r_auth = p.add_run(auth_str)
        r_auth.font.name = options.font_name
        r_auth.font.size = Pt(options.font_size_pt)

        r_year = p.add_run(f"({citation.year}). ")
        r_year.font.name = options.font_name
        r_year.font.size = Pt(options.font_size_pt)

        r_title = p.add_run(f"{citation.title}. ")
        r_title.font.name = options.font_name
        r_title.font.size = Pt(options.font_size_pt)

        if citation.journal:
            r_jour = p.add_run(f"{citation.journal}. ")
            r_jour.italic = True
            r_jour.font.name = options.font_name
            r_jour.font.size = Pt(options.font_size_pt)

        if citation.doi:
            doi_url = (
                citation.doi
                if citation.doi.startswith("http")
                else f"https://doi.org/{citation.doi}"
            )
            r_doi = p.add_run(doi_url)
            r_doi.font.name = options.font_name
            r_doi.font.size = Pt(options.font_size_pt)
        elif citation.url:
            r_url = p.add_run(citation.url)
            r_url.font.name = options.font_name
            r_url.font.size = Pt(options.font_size_pt)

    def _add_inline_formatted_runs(
        self,
        p: Paragraph,
        text: str,
        options: ExportOptionsDTO,
    ) -> None:
        """Parse markdown inline tokens (bold, italic, code) into styled runs."""
        pattern = re.compile(
            r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|___[^_]+___|__[^_]+__|_[^_]+_|`[^`]+`)"
        )
        tokens = pattern.split(text)

        for token in tokens:
            if not token:
                continue

            bold = False
            italic = False
            is_code = False
            run_text = token

            if (token.startswith("***") and token.endswith("***")) or (
                token.startswith("___") and token.endswith("___")
            ):
                bold = True
                italic = True
                run_text = token[3:-3]
            elif (token.startswith("**") and token.endswith("**")) or (
                token.startswith("__") and token.endswith("__")
            ):
                bold = True
                run_text = token[2:-2]
            elif (token.startswith("*") and token.endswith("*")) or (
                token.startswith("_") and token.endswith("_")
            ):
                italic = True
                run_text = token[1:-1]
            elif token.startswith("`") and token.endswith("`"):
                is_code = True
                run_text = token[1:-1]

            run = p.add_run(run_text)
            run.bold = bold
            run.italic = italic
            run.font.name = "Consolas" if is_code else options.font_name
            run.font.size = Pt(options.font_size_pt - 1 if is_code else options.font_size_pt)
            run.font.color.rgb = RGBColor(0, 0, 0)
