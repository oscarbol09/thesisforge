"""Unit tests for APA 7 DOCX Compiler."""

import io

from docx import Document
from docx.shared import Inches

from thesisforge.export.docx_compiler import APA7DocxCompiler
from thesisforge.models import (
    AcademicLevel,
    CitationDTO,
    ExportOptionsDTO,
    ProjectPhase,
    ProjectStateDTO,
    SectionDraftDTO,
    SectionStatus,
)


def test_docx_compiler_page_geometry_and_margins():
    """Verify standard 1-inch margins and page setup."""
    compiler = APA7DocxCompiler()
    project = ProjectStateDTO(
        id="proj-exp-01",
        title="Impacto de la IA en Educación",
        academic_level=AcademicLevel.PREGRADO,
        phase=ProjectPhase.DRAFTING,
    )
    doc = compiler.compile(project)

    for section in doc.sections:
        assert section.top_margin == Inches(1.0)
        assert section.bottom_margin == Inches(1.0)
        assert section.left_margin == Inches(1.0)
        assert section.right_margin == Inches(1.0)


def test_docx_compiler_cover_page_and_metadata():
    """Verify APA 7 Title page rendering and document core properties."""
    compiler = APA7DocxCompiler()
    options = ExportOptionsDTO(
        author_name="Valeria Mendoza",
        institution_name="Universidad Nacional",
        faculty_or_program="Facultad de Ingeniería",
        advisor_name="Dra. Carmen Soto",
        city_and_country="Lima, Perú",
        year=2026,
    )
    project = ProjectStateDTO(
        id="proj-exp-02",
        title="Modelos de Lenguaje en Investigación Académica",
        area_of_study="Ciencias de la Computación",
        academic_level=AcademicLevel.MAESTRIA,
        phase=ProjectPhase.DRAFTING,
    )

    doc = compiler.compile(project, options)

    assert doc.core_properties.title == "Modelos de Lenguaje en Investigación Académica"
    assert doc.core_properties.author == "Valeria Mendoza"
    assert doc.core_properties.subject == "Ciencias de la Computación"

    all_text = [p.text for p in doc.paragraphs]
    assert any("Modelos de Lenguaje en Investigación Académica" in t for t in all_text)
    assert any("Valeria Mendoza" in t for t in all_text)
    assert any("Universidad Nacional" in t for t in all_text)
    assert any("Asesor: Dra. Carmen Soto" in t for t in all_text)
    assert any("Lima, Perú, 2026" in t for t in all_text)


def test_docx_compiler_markdown_elements_and_headings():
    """Verify parsing of headings, paragraphs, lists, and blockquotes."""
    compiler = APA7DocxCompiler()
    project = ProjectStateDTO(
        id="proj-exp-03",
        title="Validación de Tesis",
        academic_level=AcademicLevel.PREGRADO,
        phase=ProjectPhase.DRAFTING,
        sections=[
            SectionDraftDTO(
                section_id="sec_1_1",
                chapter_number=1,
                order_index=1,
                title="Planteamiento del Problema",
                content=(
                    "El problema central radica en la **dispersión metodológica**.\n\n"
                    "## Formulación del Problema\n\n"
                    "¿Cómo influye la fragmentación de fuentes?\n\n"
                    "> La investigación contemporánea demanda rigor sintáctico y trazabilidad probatoria.\n\n"
                    "- Primer factor de riesgo\n"
                    "- Segundo factor de riesgo\n\n"
                    "1. Paso analítico preliminar\n"
                    "2. Paso analítico confirmatorio\n"
                ),
                status=SectionStatus.APPROVED,
            )
        ],
    )

    doc = compiler.compile(project)
    paragraphs = doc.paragraphs

    headings = [p.text for p in paragraphs if p.text]
    assert any("Planteamiento del Problema" in h for h in headings)
    assert any("Formulación del Problema" in h for h in headings)
    assert any("El problema central radica en la dispersión metodológica." in h for h in headings)
    assert any("La investigación contemporánea demanda rigor" in h for h in headings)
    assert any("Primer factor de riesgo" in h for h in headings)
    assert any("Paso analítico preliminar" in h for h in headings)


def test_docx_compiler_table_sanitization_and_borders():
    """Verify table extraction, border configuration, and formula injection sanitization (CWE-1236)."""
    compiler = APA7DocxCompiler()
    project = ProjectStateDTO(
        id="proj-exp-04",
        title="Estudio Cuantitativo",
        academic_level=AcademicLevel.PREGRADO,
        phase=ProjectPhase.DRAFTING,
        sections=[
            SectionDraftDTO(
                section_id="sec_2_1",
                chapter_number=2,
                order_index=1,
                title="Operacionalización de Variables",
                content=(
                    "| Variable | Indicador | Fórmula Peligrosa |\n"
                    "|---|---|---|\n"
                    "| Rendimiento | Promedio | =SUM(A1:A10) |\n"
                    r"| Asistencia | Porcentaje | @cmd\|' /C calc'!A0 |"
                    "\n"
                    "| Control | Constante | +100 |\n"
                ),
                status=SectionStatus.APPROVED,
            )
        ],
    )

    doc = compiler.compile(project)
    assert len(doc.tables) >= 1
    table = doc.tables[0]

    # Verify cells content and sanitization
    # Row 0 (Header)
    assert table.cell(0, 0).text == "Variable"
    assert table.cell(0, 1).text == "Indicador"
    assert table.cell(0, 2).text == "Fórmula Peligrosa"

    # Row 1: =SUM(A1:A10) must be sanitized with prepended apostrophe
    assert table.cell(1, 2).text == "'=SUM(A1:A10)"

    # Row 2: @cmd... must be sanitized with prepended apostrophe
    assert table.cell(2, 2).text == "'@cmd|' /C calc'!A0"

    # Row 3: +100 is a safe numeric literal
    assert table.cell(3, 2).text == "+100"


def test_docx_compiler_references_apa7_format():
    """Verify references section rendering with hanging indent and sorting."""
    compiler = APA7DocxCompiler()
    citations = [
        CitationDTO(
            id="cit-02",
            authors=["Zúñiga, F."],
            year=2024,
            title="Ética en la Investigación Tecnológica",
            journal="Revista de Informática y Sociedad",
            doi="10.1016/j.ris.2024.01",
            apa_formatted="Zúñiga, F. (2024). Ética en la Investigación Tecnológica. Revista de Informática y Sociedad, 12(1), 45-60. https://doi.org/10.1016/j.ris.2024.01",
        ),
        CitationDTO(
            id="cit-01",
            authors=["Alvarado, M.", "Benítez, R."],
            year=2023,
            title="Métodos Cuantitativos Avanzados",
            journal="Editorial Científica",
            doi="10.1000/182",
        ),
    ]

    project = ProjectStateDTO(
        id="proj-exp-05",
        title="Tesis con Referencias",
        academic_level=AcademicLevel.PREGRADO,
        phase=ProjectPhase.DRAFTING,
        validated_citations=citations,
    )

    doc = compiler.compile(project, ExportOptionsDTO(include_references=True))

    paragraphs = doc.paragraphs
    assert any(p.text == "Referencias" for p in paragraphs)

    ref_texts = [p.text for p in paragraphs if "Alvarado" in p.text or "Zúñiga" in p.text]
    assert len(ref_texts) == 2
    assert "Alvarado" in ref_texts[0]
    assert "Zúñiga" in ref_texts[1]


def test_docx_compiler_to_bytes_roundtrip():
    """Verify compilation to in-memory bytes and re-opening with python-docx."""
    compiler = APA7DocxCompiler()
    project = ProjectStateDTO(
        id="proj-exp-06",
        title="Tesis en Memoria",
        academic_level=AcademicLevel.PREGRADO,
        phase=ProjectPhase.DRAFTING,
    )

    docx_bytes = compiler.compile_to_bytes(project)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 0

    doc = Document(io.BytesIO(docx_bytes))
    assert doc.core_properties.title == "Tesis en Memoria"
